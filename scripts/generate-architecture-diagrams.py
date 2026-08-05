#!/usr/bin/env python3
"""
Enterprise Solution Architecture Diagram Generator v4
Microsoft Azure / AWS / IBM reference architecture quality.
"""
from __future__ import annotations

import html
import subprocess
import textwrap
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

OUT = Path('/opt/m-live/docs/architecture-diagrams')
W, H = 1123, 794

# Layout — target 80–85% printable area; legend in footer (outside drawing)
MARGIN = 28
HEADER_TOP = 100          # title + purpose band ends here
FOOTER_H = 44             # compact legend strip (40% smaller)
PAD = 24                  # internal box padding
GAP = 24                  # inter-box spacing
BAND_GAP = 8              # visual separation between layers
ICON = 24
BOX_H = 72
BOX_W = 168
RX = 12
STROKE_W = 2

# Typography (pt → px @ 96dpi)
FS_FIGURE = 11
FS_TITLE = 37             # 28pt
FS_PURPOSE = 19           # 14pt
FS_LAYER = 21             # 16pt bold
FS_COMPONENT = 20         # 15pt
FS_LEGEND = 15            # 11pt
FS_GROUP = 13

DIAG_W = W - 2 * MARGIN

# Enterprise palette
C = {
    'app_fill': '#DEECF9', 'app_stroke': '#0078D4', 'app_text': '#003966',
    'data_fill': '#FFF4CE', 'data_stroke': '#FFB900', 'data_text': '#5C4B00',
    'mon_fill': '#D4EDDA', 'mon_stroke': '#107C10', 'mon_text': '#0B4A0B',
    'infra_fill': '#F3F2F1', 'infra_stroke': '#605E5C', 'infra_text': '#323130',
    'ext_fill': '#E7DFEC', 'ext_stroke': '#5C2D91', 'ext_text': '#3B1F5E',
    'band': '#F5F5F5', 'band_alt': '#EBEBEB', 'band_border': '#C8C6C4', 'band_sep': '#D2D0CE',
    'white': '#FFFFFF',
    'title': '#201F1E', 'muted': '#605E5C',
    'primary_line': '#0078D4', 'secondary_line': '#8A8886',
}
FONT = 'Segoe UI, Arial, sans-serif'

Kind = Literal['app', 'data', 'mon', 'infra', 'ext']
IconKind = Literal['user', 'server', 'api', 'database', 'wallet', 'queue', 'monitor', 'shield', 'container', 'cloud']


@dataclass
class Comp:
    id: str
    label: str
    x: float
    y: float
    w: float = BOX_W
    h: float = BOX_H
    kind: Kind = 'app'
    icon: IconKind = 'server'


@dataclass
class Group:
    label: str
    x: float
    y: float
    w: float
    h: float


@dataclass
class Edge:
    src: str
    dst: str
    primary: bool = True
    label: str = ''


@dataclass
class Band:
    label: str
    y: float
    h: float
    tint: Kind | None = None


@dataclass
class Ring:
    label: str
    cx: float
    cy: float
    rx: float
    ry: float


@dataclass
class DiagramSpec:
    slug: str
    figure: str
    title: str
    purpose: str
    section: str
    insert_after: str
    description: str
    components: list[Comp] = field(default_factory=list)
    edges: list[Edge] = field(default_factory=list)
    bands: list[Band] = field(default_factory=list)
    groups: list[Group] = field(default_factory=list)
    rings: list[Ring] = field(default_factory=list)
    pipeline: list[tuple[str, Kind, IconKind]] = field(default_factory=list)
    pipeline_x: float = 0
    pipeline_y: float = 0
    pipeline_side: list[Comp] = field(default_factory=list)


# Fluent UI / Material-style outline icons (single family)
ICON_PATHS = {
    'user': 'M10 2a4 4 0 1 0 0 8 4 4 0 0 0 0-8zm-8 12c0-3.3 3.6-6 8-6s8 2.7 8 6v1H2v-1z',
    'server': 'M2 3h16v5H2V3zm0 7h16v5H2v-5zm3 1.5h2v1H5v-1zm0-6h2v1H5v-1z',
    'api': 'M3 5h14v2H3V5zm0 4h10v2H3V9zm0 4h12v2H3v-2z',
    'database': 'M10 2c4.4 0 8 1.8 8 4s-3.6 4-8 4-8-1.8-8-4 3.6-4 8-4zm-8 6v2c0 2.2 3.6 4 8 4s8-1.8 8-4V8M2 12v2c0 2.2 3.6 4 8 4s8-1.8 8-4v-2',
    'wallet': 'M3 6h14a2 2 0 0 1 2 2v1H3V6zm0 3h16v5a2 2 0 0 1-2 2H3a2 2 0 0 1-2-2V9h16V7H3z',
    'queue': 'M3 3h14v4H3V3zm0 6h14v4H3V9zm0 6h10v4H3v-4z',
    'monitor': 'M2 3h16v9H2V3zm5 10h6v3H7v-3z',
    'shield': 'M10 1l8 3.5v4.5c0 5.2-3.8 8.8-8 9.5C5.8 17.8 2 14.2 2 9V4.5L10 1z',
    'container': 'M2 5h16v10H2V5zm3 2v2h2V7H5zm5 0v2h2V7h-2zm5 0v2h2V7h-2z',
    'cloud': 'M5 14h10a3.5 3.5 0 0 0 .4-7A4.5 4.5 0 0 0 6.5 4.5 3.5 3.5 0 0 0 5 14z',
}


def colors(kind: Kind) -> tuple[str, str, str]:
    return C[f'{kind}_fill'], C[f'{kind}_stroke'], C[f'{kind}_text']


def comp_map(spec: DiagramSpec) -> dict[str, Comp]:
    m = {c.id: c for c in spec.components}
    for i, (lab, k, ic) in enumerate(spec.pipeline):
        cid = f'pl{i}'
        m[cid] = Comp(cid, lab, 0, 0, kind=k, icon=ic)
    for c in spec.pipeline_side:
        m[c.id] = c
    return m


def layout_pipeline(spec: DiagramSpec, max_per_row: int = 5) -> None:
    if not spec.pipeline:
        return
    items = spec.pipeline
    rows = [items[i:i + max_per_row] for i in range(0, len(items), max_per_row)]
    x0, y0 = spec.pipeline_x, spec.pipeline_y
    avail = DIAG_W - 32
    for ri, row in enumerate(rows):
        n = len(row)
        bw = min(BOX_W + 32, (avail - (n - 1) * GAP) / max(n, 1))
        for ci, (lab, k, ic) in enumerate(row):
            idx = sum(len(r) for r in rows[:ri]) + ci
            cid = f'pl{idx}'
            spec.components.append(Comp(
                cid, lab, x0 + ci * (bw + GAP), y0 + ri * (BOX_H + GAP + 20),
                w=bw, h=BOX_H, kind=k, icon=ic,
            ))
    ordered = [f'pl{i}' for i in range(len(items))]
    for i in range(len(ordered) - 1):
        spec.edges.append(Edge(ordered[i], ordered[i + 1], primary=True))


def layout_bands_stack(bands_labels: list[str], top: float, bottom: float, tints: list[Kind | None] | None = None) -> list[Band]:
    n = len(bands_labels)
    total_gap = BAND_GAP * max(n - 1, 0)
    h = (bottom - top - total_gap) / n
    bands = []
    y = top
    for i, lab in enumerate(bands_labels):
        tint = tints[i] if tints and i < len(tints) else None
        bands.append(Band(lab, y, h, tint=tint))
        y += h + BAND_GAP
    return bands


def place_row_in_band(comps: list[tuple[str, Kind, IconKind]], band: Band, start_x: float | None = None) -> list[Comp]:
    n = len(comps)
    sx = start_x if start_x is not None else MARGIN + 20
    avail = MARGIN + DIAG_W - sx
    bw = min(BOX_W + 28, (avail - (n - 1) * GAP) / max(n, 1))
    cy = band.y + (band.h - BOX_H) / 2
    out = []
    slug = ''.join(c for c in band.label[:4] if c.isalnum()) or 'row'
    for i, (lab, k, ic) in enumerate(comps):
        out.append(Comp(f'b{slug}{i}', lab, sx + i * (bw + GAP), cy, w=bw, h=BOX_H, kind=k, icon=ic))
    return out


def svg_defs() -> str:
    return f'''<defs>
  <filter id="shadow" x="-4%" y="-4%" width="108%" height="108%">
    <feDropShadow dx="0" dy="2" stdDeviation="3" flood-color="#000000" flood-opacity="0.10"/>
  </filter>
  <marker id="pa" markerWidth="10" markerHeight="7" refX="9" refY="3.5" orient="auto">
    <polygon points="0,0 10,3.5 0,7" fill="{C['primary_line']}"/>
  </marker>
  <marker id="sa" markerWidth="8" markerHeight="6" refX="7" refY="3" orient="auto">
    <polygon points="0,0 8,3 0,6" fill="{C['secondary_line']}"/>
  </marker>
  {''.join(f'<symbol id="ic-{n}" viewBox="0 0 20 18"><path d="{d}" fill="currentColor"/></symbol>' for n,d in ICON_PATHS.items())}
</defs>'''


def render_icon(kind: IconKind, x: float, y: float, stroke: str) -> str:
    return f'<use href="#ic-{kind}" x="{x:.0f}" y="{y:.0f}" width="{ICON}" height="{ICON}" color="{stroke}"/>'


def wrap_label(label: str, max_chars: int = 14) -> list[str]:
    lines = []
    for part in label.replace('\n', ' ').split(' / '):
        if len(part) <= max_chars:
            lines.append(part)
        else:
            lines.extend(textwrap.wrap(part, width=max_chars))
    return lines[:2]


def render_svg(spec: DiagramSpec) -> str:
    layout_pipeline(spec)
    cm = comp_map(spec)
    diag_bottom = H - FOOTER_H - 8

    parts = [
        f'<?xml version="1.0" encoding="UTF-8"?>\n',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">\n',
        svg_defs(),
        f'<rect width="{W}" height="{H}" fill="{C["white"]}"/>\n',
        f'<text x="{MARGIN}" y="28" font-family="{FONT}" font-size="{FS_FIGURE}" fill="{C["muted"]}">{html.escape(spec.figure)}</text>\n',
        f'<text x="{MARGIN}" y="58" font-family="{FONT}" font-size="{FS_TITLE}" font-weight="700" fill="{C["title"]}">{html.escape(spec.title)}</text>\n',
        f'<text x="{MARGIN}" y="82" font-family="{FONT}" font-size="{FS_PURPOSE}" fill="{C["muted"]}"><tspan font-weight="700">Purpose: </tspan>{html.escape(spec.purpose)}</text>\n',
        f'<line x1="{MARGIN}" y1="90" x2="{W-MARGIN}" y2="90" stroke="{C["band_sep"]}" stroke-width="1"/>\n',
    ]

    for i, b in enumerate(spec.bands):
        if b.tint:
            fill, stroke, _ = colors(b.tint)
            fill = fill
            border = stroke
        else:
            fill = C['band'] if i % 2 == 0 else C['band_alt']
            border = C['band_border']
        parts.append(
            f'<rect x="{MARGIN}" y="{b.y:.0f}" width="{DIAG_W:.0f}" height="{b.h:.0f}" fill="{fill}" fill-opacity="0.55" stroke="{border}" stroke-width="1.5"/>\n'
            f'<text x="{MARGIN+14}" y="{b.y+26:.0f}" font-family="{FONT}" font-size="{FS_LAYER}" font-weight="700" fill="{C["title"]}">{html.escape(b.label)}</text>\n'
        )
        if i < len(spec.bands) - 1:
            ny = b.y + b.h + BAND_GAP / 2
            parts.append(f'<line x1="{MARGIN}" y1="{ny:.0f}" x2="{MARGIN+DIAG_W:.0f}" y2="{ny:.0f}" stroke="{C["band_sep"]}" stroke-width="1"/>\n')

    cx_full = MARGIN + DIAG_W / 2
    cy_full = HEADER_TOP + (diag_bottom - HEADER_TOP) / 2
    for ring in spec.rings:
        fill, stroke, _ = colors('infra')
        parts.append(
            f'<ellipse cx="{cx_full:.0f}" cy="{cy_full:.0f}" rx="{ring.rx:.0f}" ry="{ring.ry:.0f}" fill="{fill}" fill-opacity="0.65" stroke="{stroke}" stroke-width="{STROKE_W}"/>\n'
            f'<text x="{cx_full:.0f}" y="{cy_full-ring.ry+22:.0f}" text-anchor="middle" font-family="{FONT}" font-size="{FS_COMPONENT}" font-weight="700" fill="{C["title"]}">{html.escape(ring.label)}</text>\n'
        )

    for g in spec.groups:
        parts.append(
            f'<rect x="{g.x:.0f}" y="{g.y:.0f}" width="{g.w:.0f}" height="{g.h:.0f}" rx="{RX}" fill="none" stroke="{C["app_stroke"]}" stroke-width="1.5" stroke-dasharray="8 5"/>\n'
            f'<text x="{g.x+12:.0f}" y="{g.y+18:.0f}" font-family="{FONT}" font-size="{FS_GROUP}" font-weight="700" fill="{C["app_stroke"]}">{html.escape(g.label)}</text>\n'
        )

    for e in spec.edges:
        if e.src not in cm or e.dst not in cm:
            continue
        s, t = cm[e.src], cm[e.dst]
        x1, y1 = s.x + s.w / 2, s.y + s.h / 2
        x2, y2 = t.x + t.w / 2, t.y + t.h / 2
        if abs(y1 - y2) < 10 and s.x + s.w <= t.x:
            path = f'M {s.x+s.w:.0f} {y1:.0f} L {t.x:.0f} {y2:.0f}'
        elif y1 < y2 - 10:
            mid = (s.y + s.h + t.y) / 2
            path = f'M {x1:.0f} {s.y+s.h:.0f} L {x1:.0f} {mid:.0f} L {x2:.0f} {mid:.0f} L {x2:.0f} {t.y:.0f}'
        elif s.x + s.w <= t.x:
            path = f'M {s.x+s.w:.0f} {y1:.0f} L {t.x:.0f} {y2:.0f}'
        else:
            mid = (s.y + s.h + t.y) / 2
            path = f'M {x1:.0f} {s.y+s.h:.0f} L {x1:.0f} {mid:.0f} L {x2:.0f} {mid:.0f} L {x2:.0f} {t.y:.0f}'
        sw = '2.5' if e.primary else '1'
        col = C['primary_line'] if e.primary else C['secondary_line']
        mk = 'pa' if e.primary else 'sa'
        dash = '' if e.primary else ' stroke-dasharray="5 4"'
        parts.append(f'<path d="{path}" fill="none" stroke="{col}" stroke-width="{sw}"{dash} marker-end="url(#{mk})"/>\n')

    for c in spec.components:
        fill, stroke, tcol = colors(c.kind)
        parts.append(f'<g id="{html.escape(c.id)}">\n')
        parts.append(
            f'  <rect x="{c.x:.0f}" y="{c.y:.0f}" width="{c.w:.0f}" height="{c.h:.0f}" rx="{RX}" fill="{fill}" stroke="{stroke}" stroke-width="{STROKE_W}" filter="url(#shadow)"/>\n'
        )
        iy = c.y + (c.h - ICON) / 2
        parts.append(f'  <g transform="translate({c.x+PAD:.0f},{iy:.0f})">{render_icon(c.icon, 0, 0, stroke)}</g>\n')
        lines = wrap_label(c.label, 16)
        tx = c.x + PAD + ICON + 8
        ty = c.y + c.h / 2 + 5 - (len(lines) - 1) * 8
        for i, line in enumerate(lines):
            parts.append(
                f'  <text x="{tx:.0f}" y="{ty + i*16:.0f}" font-family="{FONT}" font-size="{FS_COMPONENT}" font-weight="600" fill="{tcol}">{html.escape(line)}</text>\n'
            )
        parts.append('</g>\n')

    # Footer legend — outside drawing area
    ly = H - FOOTER_H + 6
    parts.append(f'<line x1="{MARGIN}" y1="{ly-4:.0f}" x2="{W-MARGIN}" y2="{ly-4:.0f}" stroke="{C["band_sep"]}" stroke-width="1"/>\n')
    lx = MARGIN + 8
    parts.append(f'<text x="{lx:.0f}" y="{ly+14:.0f}" font-family="{FONT}" font-size="{FS_LEGEND}" font-weight="700" fill="{C["title"]}">Legend:</text>\n')
    xoff = lx + 58
    for lab, k in [('Application', 'app'), ('Data Store', 'data'), ('Monitoring', 'mon'), ('Infrastructure', 'infra'), ('External', 'ext')]:
        fill, stroke, _ = colors(k)
        parts.append(f'<rect x="{xoff:.0f}" y="{ly+2:.0f}" width="16" height="12" rx="3" fill="{fill}" stroke="{stroke}" stroke-width="1.5"/>\n')
        parts.append(f'<text x="{xoff+22:.0f}" y="{ly+14:.0f}" font-family="{FONT}" font-size="{FS_LEGEND}" fill="{C["muted"]}">{lab}</text>\n')
        xoff += 108
    parts.append(f'<line x1="{xoff:.0f}" y1="{ly+8:.0f}" x2="{xoff+22:.0f}" y2="{ly+8:.0f}" stroke="{C["primary_line"]}" stroke-width="2.5" marker-end="url(#pa)"/>\n')
    parts.append(f'<text x="{xoff+28:.0f}" y="{ly+14:.0f}" font-family="{FONT}" font-size="{FS_LEGEND}" fill="{C["muted"]}">Primary</text>\n')
    xoff += 88
    parts.append(f'<line x1="{xoff:.0f}" y1="{ly+8:.0f}" x2="{xoff+22:.0f}" y2="{ly+8:.0f}" stroke="{C["secondary_line"]}" stroke-width="1" stroke-dasharray="5 4" marker-end="url(#sa)"/>\n')
    parts.append(f'<text x="{xoff+28:.0f}" y="{ly+14:.0f}" font-family="{FONT}" font-size="{FS_LEGEND}" fill="{C["muted"]}">Dependency</text>\n')

    parts.append('</svg>\n')
    return ''.join(parts)


def render_drawio(spec: DiagramSpec) -> str:
    layout_pipeline(spec)
    root = ET.Element('mxfile', host='app.diagrams.net', agent='enterprise-v4')
    d_el = ET.SubElement(root, 'diagram', name=spec.title, id=spec.slug)
    model = ET.SubElement(d_el, 'mxGraphModel', pageWidth='1169', pageHeight='827', grid='1')
    rc = ET.SubElement(model, 'root')
    ET.SubElement(rc, 'mxCell', id='0')
    ET.SubElement(rc, 'mxCell', id='1', parent='0')
    nid = 2

    def add(val, st, x, y, w, h, cid=None, edge=False, src=None, dst=None):
        nonlocal nid
        i = cid or str(nid)
        nid += 1
        a = {'id': i, 'value': val.replace('\n', '<br>'), 'style': st, 'parent': '1'}
        if edge:
            a.update(edge='1', source=src, target=dst)
        else:
            a['vertex'] = '1'
        c = ET.SubElement(rc, 'mxCell', **a)
        g = ET.SubElement(c, 'mxGeometry', x=str(x), y=str(y), width=str(w), height=str(h))
        g.set('as', 'geometry')
        return i

    add(f'{spec.figure} — {spec.title}', f'text;fontSize={FS_TITLE};fontStyle=1;fontColor=#201F1E;fontFamily=Segoe UI;', MARGIN, 16, 820, 44, cid='title')
    add(f'Purpose: {spec.purpose}', f'text;fontSize={FS_PURPOSE};fontColor=#605E5C;fontFamily=Segoe UI;', MARGIN, 62, 900, 28, cid='purpose')

    for b in spec.bands:
        fill = '#F5F5F5' if not b.tint else colors(b.tint)[0]
        add('', f'rounded=1;fillColor={fill};strokeColor=#C8C6C4;strokeWidth=2;', MARGIN, b.y, DIAG_W, b.h)
        add(b.label, f'text;fontSize={FS_LAYER};fontStyle=1;fontColor=#201F1E;fontFamily=Segoe UI;', MARGIN + 14, b.y + 8, 320, 22)

    ids = {}
    for c in spec.components:
        fill, stroke, _ = colors(c.kind)
        st = f'rounded=1;arcSize=12;whiteSpace=wrap;html=1;fillColor={fill};strokeColor={stroke};strokeWidth=2;fontStyle=1;fontSize={FS_COMPONENT};fontFamily=Segoe UI;shadow=1;'
        ids[c.id] = add(c.label, st, c.x, c.y, c.w, c.h, cid=c.id)

    for e in spec.edges:
        if e.src in ids and e.dst in ids:
            st = 'edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;strokeColor=#0078D4;strokeWidth=2.5;endArrow=block;' if e.primary else 'edgeStyle=orthogonalEdgeStyle;dashed=1;strokeColor=#8A8886;strokeWidth=1;endArrow=block;'
            add('', st, 0, 0, 0, 0, edge=True, src=ids[e.src], dst=ids[e.dst])

    return '<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(root, encoding='unicode')


def render_pptx(spec: DiagramSpec, path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    layout_pipeline(spec)
    prs = Presentation()
    prs.slide_width = Inches(11.69)
    prs.slide_height = Inches(8.27)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def rgb(h: str) -> RGBColor:
        h = h.lstrip('#')
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    tb = slide.shapes.add_textbox(Inches(0.35), Inches(0.08), Inches(10.5), Inches(0.45))
    tb.text_frame.text = f'{spec.figure}  {spec.title}'
    tb.text_frame.paragraphs[0].font.size = Pt(28)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.name = 'Segoe UI'

    pb = slide.shapes.add_textbox(Inches(0.35), Inches(0.48), Inches(10.5), Inches(0.32))
    pb.text_frame.text = f'Purpose: {spec.purpose}'
    pb.text_frame.paragraphs[0].font.size = Pt(14)
    pb.text_frame.paragraphs[0].font.name = 'Segoe UI'

    for b in spec.bands:
        fill_hex = C['band'] if not b.tint else colors(b.tint)[0]
        sh = slide.shapes.add_shape(1, Inches(MARGIN / 96), Inches(b.y / 96), Inches(DIAG_W / 96), Inches(b.h / 96))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill_hex)
        sh.line.color.rgb = rgb(C['band_border'])
        sh.line.width = Pt(1.5)

    for c in spec.components:
        fill, stroke, _ = colors(c.kind)
        sh = slide.shapes.add_shape(1, Inches(c.x / 96), Inches(c.y / 96), Inches(c.w / 96), Inches(c.h / 96))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        sh.line.color.rgb = rgb(stroke)
        sh.line.width = Pt(2)
        sh.text_frame.text = c.label
        sh.text_frame.margin_left = Inches(PAD / 96)
        for p in sh.text_frame.paragraphs:
            p.font.size = Pt(15)
            p.font.name = 'Segoe UI'
            p.font.bold = True
            p.alignment = PP_ALIGN.LEFT

    prs.save(str(path))


def centered_pipeline_y(n_items: int, max_per_row: int = 5) -> float:
    rows = max(1, (n_items + max_per_row - 1) // max_per_row)
    block_h = rows * BOX_H + max(0, rows - 1) * (GAP + 20)
    return HEADER_TOP + (H - FOOTER_H - HEADER_TOP - block_h) / 2


def export_png(svg: Path, png: Path) -> None:
    subprocess.run(['rsvg-convert', '-w', '3508', '-h', '2480', str(svg), '-o', str(png)], check=True)


def export_pdf(svg: Path, pdf: Path) -> None:
    subprocess.run(['rsvg-convert', '-w', '3508', '-h', '2480', str(svg), '-f', 'pdf', '-o', str(pdf)], check=True)


def build_all() -> list[DiagramSpec]:
    top = HEADER_TOP + 4
    bottom = H - FOOTER_H - 12
    specs: list[DiagramSpec] = []

    # ── 2.1 Enterprise Solution Architecture ──
    bands = layout_bands_stack([
        'Client Layer', 'Edge Layer', 'Application Layer', 'Financial Layer',
        'Data Layer', 'Observability Layer', 'Infrastructure Layer',
    ], top, bottom)
    s = DiagramSpec(
        slug='01-Enterprise-Solution-Architecture', figure='Figure 2.1',
        title='Enterprise Solution Architecture',
        purpose='How users, edge services, applications, financial engines, data stores, and observability compose the production deployment.',
        section='2. Solution Architecture', insert_after='Section 2 Introduction',
        description='Seven-layer deployment topology for METHErium CEX on VPS Docker Compose.',
        bands=bands,
    )
    layer_comps = [
        [('Retail Users', 'ext', 'user'), ('Institutional Users', 'ext', 'user'), ('Administrators', 'ext', 'user'), ('Mobile Clients', 'ext', 'user')],
        [('Internet', 'ext', 'cloud'), ('Reverse Proxy', 'infra', 'server'), ('TLS Termination', 'infra', 'shield'), ('Load Distribution', 'infra', 'server')],
        [('Web Application', 'app', 'server'), ('Administration Portal', 'app', 'server'), ('API Gateway', 'app', 'api'), ('Matching Engine', 'app', 'server')],
        [('Settlement Service', 'app', 'server'), ('Wallet Service', 'app', 'wallet'), ('Treasury Service', 'app', 'wallet'), ('Indexer Service', 'app', 'server')],
        [('Operational Database', 'data', 'database'), ('Distributed Cache', 'data', 'database'), ('Message Broker', 'data', 'queue'), ('Event Streaming Platform', 'data', 'queue')],
        [('Metrics Collection', 'mon', 'monitor'), ('Dashboards', 'mon', 'monitor'), ('Health Probes', 'mon', 'monitor'), ('Alert Rules', 'mon', 'monitor')],
        [('Container Runtime', 'infra', 'container'), ('Private Network', 'infra', 'server'), ('Key Management', 'ext', 'shield'), ('External RPC', 'ext', 'cloud')],
    ]
    for band, row in zip(bands, layer_comps):
        s.components.extend(place_row_in_band(row, band))
    for i in range(len(bands) - 1):
        a = s.components[i * 4]
        b = s.components[(i + 1) * 4]
        s.edges.append(Edge(a.id, b.id, primary=True))
    specs.append(s)

    # ── 2.2 Logical Layered Architecture ──
    bands2 = layout_bands_stack([
        'Presentation Layer', 'Application Layer', 'Business Layer',
        'Financial Layer', 'Infrastructure Layer', 'Observability Layer',
    ], top, bottom)
    s = DiagramSpec(
        slug='02-Logical-Layered-Architecture', figure='Figure 2.2',
        title='Logical Layered Architecture',
        purpose='Logical separation of presentation, application, business, financial, infrastructure, and observability concerns.',
        section='2. Solution Architecture', insert_after='Figure 2.1',
        description='Six distinct logical layers without merging business and application tiers.',
        bands=bands2,
    )
    rows2 = [
        [('Web Experience', 'app', 'server'), ('Admin Experience', 'app', 'server'), ('Mobile Experience', 'app', 'server'), ('Public API Docs', 'app', 'api')],
        [('API Platform', 'app', 'api'), ('Matching Platform', 'app', 'server'), ('Integration Hub', 'app', 'api'), ('Session Platform', 'app', 'shield')],
        [('Identity & Access', 'app', 'shield'), ('Spot Trading', 'app', 'server'), ('P2P Marketplace', 'app', 'server'), ('Compliance Controls', 'app', 'shield')],
        [('Settlement Engine', 'app', 'server'), ('Wallet & Ledger', 'app', 'wallet'), ('Treasury Operations', 'app', 'wallet'), ('Fee & Risk Controls', 'app', 'shield')],
        [('Relational Store', 'data', 'database'), ('Cache & Sessions', 'data', 'database'), ('Async Messaging', 'data', 'queue'), ('Stream Processing', 'data', 'queue')],
        [('Metrics Pipeline', 'mon', 'monitor'), ('Operational Dashboards', 'mon', 'monitor'), ('SLO Monitoring', 'mon', 'monitor'), ('Incident Signals', 'mon', 'monitor')],
    ]
    for band, row in zip(bands2, rows2):
        s.components.extend(place_row_in_band(row, band))
    specs.append(s)

    # ── 2.3 Component Interaction ──
    s = DiagramSpec(
        slug='03-Component-Interaction-Diagram', figure='Figure 2.3',
        title='Component Interaction Diagram',
        purpose='Which major platform modules interact to deliver exchange capabilities.',
        section='2. Solution Architecture', insert_after='Figure 2.2',
        description='Module-level interaction map across authentication, trading, wallet, settlement, treasury, compliance, notifications, monitoring, and admin.',
        bands=layout_bands_stack(['Experience & Access', 'Core Platform Modules', 'Shared Platform Services'], top, bottom),
    )
    b0, b1, b2 = s.bands
    s.components.extend(place_row_in_band([
        ('Authentication', 'app', 'shield'), ('User Management', 'app', 'user'), ('Admin Operations', 'app', 'server'),
    ], b0))
    s.components.extend(place_row_in_band([
        ('Trading', 'app', 'server'), ('Wallet', 'app', 'wallet'), ('Settlement', 'app', 'server'), ('Treasury', 'app', 'wallet'), ('P2P', 'app', 'server'),
    ], b1))
    s.components.extend(place_row_in_band([
        ('Compliance', 'app', 'shield'), ('Notifications', 'app', 'queue'), ('Monitoring', 'mon', 'monitor'), ('Operational Database', 'data', 'database'),
    ], b2))
    ids = [c.id for c in s.components]
    s.edges += [Edge(ids[0], ids[1], True), Edge(ids[3], ids[6], True), Edge(ids[4], ids[6], True), Edge(ids[5], ids[7], True),
                Edge(ids[6], ids[10], False), Edge(ids[8], ids[11], False), Edge(ids[9], ids[11], False)]
    specs.append(s)

    # ── 4.1 Authentication ──
    s = DiagramSpec(
        slug='04-Authentication-Flow', figure='Figure 4.1',
        title='Authentication Architecture',
        purpose='How registration, login, token issuance, session management, and role-based access enable secure platform entry.',
        section='4. Application Architecture', insert_after='Section 4 Introduction',
        description='Authentication flow from registration through dashboard access.',
        pipeline=[
            ('Registration', 'app', 'user'), ('Login', 'app', 'shield'), ('Credential Validation', 'app', 'shield'),
            ('Multi-Factor Auth', 'app', 'shield'), ('Token Issuance', 'app', 'api'), ('Session Store', 'data', 'database'),
            ('Role-Based Access', 'app', 'shield'), ('Secure Dashboard', 'app', 'server'),
        ],
        pipeline_x=MARGIN + 20, pipeline_y=centered_pipeline_y(8),
    )
    specs.append(s)

    # ── 4.2 Trading Lifecycle ──
    s = DiagramSpec(
        slug='05-Trading-Lifecycle', figure='Figure 4.2',
        title='Trading Lifecycle Architecture',
        purpose='End-to-end spot order processing from submission through balance update and client notification.',
        section='4. Application Architecture', insert_after='Figure 4.1',
        description='Trading lifecycle across validation, matching, settlement, ledger, wallet, and notification.',
        pipeline=[
            ('Order Submission', 'app', 'api'), ('Risk Validation', 'app', 'shield'), ('Matching Engine', 'app', 'server'),
            ('Trade Capture', 'app', 'server'), ('Settlement', 'app', 'server'), ('Ledger Posting', 'data', 'database'),
            ('Balance Update', 'data', 'wallet'), ('Client Notification', 'app', 'queue'),
        ],
        pipeline_x=MARGIN + 10, pipeline_y=centered_pipeline_y(8),
    )
    specs.append(s)

    # ── 4.3 Matching Engine ──
    s = DiagramSpec(
        slug='06-Matching-Engine', figure='Figure 4.3',
        title='Matching Engine Architecture',
        purpose='How orders are validated, matched, persisted, and published as trade events for downstream settlement.',
        section='4. Application Architecture', insert_after='Figure 4.2',
        description='Matching pipeline with durability, event bus, and output interfaces.',
        bands=layout_bands_stack(['Order Processing Pipeline', 'Durability & Event Distribution'], top, bottom),
    )
    proc, dur = s.bands
    s.components.extend(place_row_in_band([
        ('Order Input', 'app', 'api'), ('Validation', 'app', 'shield'), ('Order Book', 'data', 'database'),
        ('Price-Time Priority', 'app', 'server'), ('Matching Engine', 'app', 'server'), ('Trade Generation', 'app', 'server'),
    ], proc))
    s.components.extend(place_row_in_band([
        ('Write-Ahead Log', 'data', 'database'), ('Event Streaming Platform', 'data', 'queue'), ('Event Bus', 'data', 'queue'),
        ('Output APIs', 'app', 'api'), ('Settlement Handoff', 'app', 'server'),
    ], dur))
    proc_ids = [c.id for c in s.components[:6]]
    dur_ids = [c.id for c in s.components[6:]]
    for i in range(len(proc_ids) - 1):
        s.edges.append(Edge(proc_ids[i], proc_ids[i + 1], primary=True))
    for i in range(len(dur_ids) - 1):
        s.edges.append(Edge(dur_ids[i], dur_ids[i + 1], primary=True))
    s.edges += [
        Edge(proc_ids[-1], dur_ids[0], primary=True),
        Edge(proc_ids[4], dur_ids[0], primary=False),
        Edge(dur_ids[1], dur_ids[2], primary=False),
    ]
    specs.append(s)

    # ── 4.4 Settlement ──
    s = DiagramSpec(
        slug='07-Settlement-Engine', figure='Figure 4.4',
        title='Settlement Architecture',
        purpose='How trade events become ledger entries, wallet balance updates, audit records, and user notifications.',
        section='4. Application Architecture', insert_after='Figure 4.3',
        description='Settlement pipeline with queue, validation, ledger, balances, audit, and notification.',
        pipeline=[
            ('Trade Event', 'app', 'queue'), ('Settlement Queue', 'data', 'queue'), ('Validation', 'app', 'shield'),
            ('Ledger', 'data', 'database'), ('Wallet Balance', 'data', 'wallet'), ('Audit Storage', 'data', 'database'),
            ('Notification', 'app', 'queue'),
        ],
        pipeline_x=MARGIN + 40, pipeline_y=centered_pipeline_y(7),
    )
    s.groups = [Group('Financial Integrity Controls', MARGIN + 20, top + 260, DIAG_W - 40, 90)]
    specs.append(s)

    # ── 4.5 Wallet ──
    s = DiagramSpec(
        slug='08-Wallet-Architecture', figure='Figure 4.5',
        title='Wallet Architecture',
        purpose='Deposit and withdrawal processing including indexing, signing, chain connectivity, and balance management.',
        section='4. Application Architecture', insert_after='Figure 4.4',
        description='Wallet architecture with separate deposit and withdrawal flows.',
        bands=layout_bands_stack(
            ['Deposit Pipeline', 'Withdrawal Pipeline', 'Wallet Platform Services'],
            top, bottom, tints=['app', 'ext', None],
        ),
    )
    dep, wit, plat = s.bands
    s.components.extend(place_row_in_band([
        ('Deposit Request', 'app', 'api'), ('Blockchain Indexer', 'app', 'server'), ('Credit Processing', 'app', 'wallet'),
    ], dep))
    s.components.extend(place_row_in_band([
        ('Withdrawal Request', 'app', 'api'), ('Signing Service', 'app', 'shield'), ('Chain Broadcast', 'ext', 'cloud'),
    ], wit))
    s.components.extend(place_row_in_band([
        ('Hot Wallet', 'data', 'wallet'), ('Cold Wallet', 'data', 'wallet'), ('Ledger', 'data', 'database'), ('Balances', 'data', 'database'),
    ], plat))
    s.groups = [
        Group('Deposit Pipeline', MARGIN + 12, dep.y + 22, DIAG_W - 24, dep.h - 28),
        Group('Withdrawal Pipeline', MARGIN + 12, wit.y + 22, DIAG_W - 24, wit.h - 28),
    ]
    dep_ids = [c.id for c in s.components[:3]]
    wit_ids = [c.id for c in s.components[3:6]]
    for ids in (dep_ids, wit_ids):
        for i in range(len(ids) - 1):
            s.edges.append(Edge(ids[i], ids[i + 1], primary=True))
    specs.append(s)

    # ── 4.6 Treasury ──
    s = DiagramSpec(
        slug='09-Treasury-Architecture', figure='Figure 4.6',
        title='Treasury Architecture',
        purpose='Treasury operations spanning hot and cold wallets, sweeps, reconciliation, key management, audit, and risk controls.',
        section='4. Application Architecture', insert_after='Figure 4.5',
        description='Treasury control plane and wallet stores.',
        bands=layout_bands_stack(['Wallet Stores', 'Treasury Operations', 'Controls & Audit'], top, bottom),
    )
    wband, oband, cband = s.bands
    s.components.extend(place_row_in_band([('Hot Wallet Store', 'data', 'database'), ('Cold Wallet Store', 'data', 'database'), ('Reserve Policy', 'app', 'shield')], wband))
    s.components.extend(place_row_in_band([
        ('Treasury Operations', 'app', 'wallet'), ('Deposit Sweeps', 'app', 'server'), ('Reconciliation', 'app', 'monitor'),
    ], oband))
    s.components.extend(place_row_in_band([
        ('Key Management', 'ext', 'shield'), ('Audit Trail', 'data', 'database'), ('Risk Controls', 'app', 'shield'),
    ], cband))
    specs.append(s)

    # ── 4.7 Market Data ──
    s = DiagramSpec(
        slug='10-Market-Data-Architecture', figure='Figure 4.7',
        title='Market Data Architecture',
        purpose='How trades, tickers, charts, and order books are served via REST and WebSocket channels.',
        section='4. Application Architecture', insert_after='Figure 4.6',
        description='Market data sources, caching, streaming, and delivery interfaces.',
        bands=layout_bands_stack(['Market Data Sources', 'Distribution Services', 'Client Channels'], top, bottom),
    )
    for band, row in zip(s.bands, [
        [('Trade Feed', 'data', 'database'), ('Ticker Service', 'app', 'server'), ('Chart Aggregates', 'data', 'database'), ('Order Book', 'data', 'database')],
        [('Stream Processor', 'app', 'server'), ('Distributed Cache', 'data', 'database'), ('Event Streaming Platform', 'data', 'queue')],
        [('REST API', 'app', 'api'), ('WebSocket Feed', 'app', 'queue'), ('Market Snapshots', 'app', 'api')],
    ]):
        s.components.extend(place_row_in_band(row, band))
    specs.append(s)

    # ── 4.8 P2P ──
    s = DiagramSpec(
        slug='11-P2P-Escrow-Flow', figure='Figure 4.8',
        title='P2P Escrow Architecture',
        purpose='Peer-to-peer trade lifecycle including escrow, payment verification, release, dispute, admin review, and settlement.',
        section='4. Application Architecture', insert_after='Figure 4.7',
        description='Full P2P escrow lifecycle for buyer, seller, admin, and platform services.',
        pipeline=[
            ('Buyer', 'ext', 'user'), ('Seller', 'ext', 'user'), ('Escrow Lock', 'app', 'wallet'), ('Payment', 'app', 'api'),
            ('Verification', 'app', 'shield'), ('Release', 'app', 'wallet'), ('Dispute', 'app', 'shield'),
            ('Admin Review', 'app', 'server'), ('Settlement', 'app', 'server'), ('Notifications', 'app', 'queue'),
        ],
        pipeline_x=MARGIN + 8, pipeline_y=centered_pipeline_y(10),
    )
    # layout_pipeline adds sequential edges pl0→pl1→…→pl9 across both rows
    specs.append(s)

    # ── 5.1 Defense in Depth ──
    s = DiagramSpec(
        slug='12-Defense-in-Depth', figure='Figure 5.1',
        title='Defense in Depth Architecture',
        purpose='Layered security controls from network perimeter through application, financial processing, and immutable audit.',
        section='5. Security Architecture', insert_after='Section 5 Introduction',
        description='Sequential security layers protecting the exchange platform.',
        pipeline=[
            ('Internet Perimeter', 'ext', 'cloud'), ('Edge Protection', 'infra', 'shield'), ('Authentication', 'app', 'shield'),
            ('Authorization', 'app', 'shield'), ('Business Validation', 'app', 'server'), ('Financial Controls', 'app', 'wallet'),
            ('Data Protection', 'data', 'database'), ('Immutable Audit', 'data', 'database'),
        ],
        pipeline_x=MARGIN + 10, pipeline_y=centered_pipeline_y(8),
    )
    specs.append(s)

    # ── 5.2 Security Layers (concentric) ──
    cx, cy = MARGIN + DIAG_W / 2, HEADER_TOP + (bottom - HEADER_TOP) / 2
    s = DiagramSpec(
        slug='13-Security-Layers', figure='Figure 5.2',
        title='Security Layer Model',
        purpose='Concentric security domains from network perimeter to financial core and audit.',
        section='5. Security Architecture', insert_after='Figure 5.1',
        description='Concentric security layers: network, infrastructure, application, wallet, financial, audit.',
        rings=[
            Ring('Network Security', cx, cy, 340, 200),
            Ring('Infrastructure Security', cx, cy, 280, 165),
            Ring('Application Security', cx, cy, 220, 130),
            Ring('Wallet Security', cx, cy, 160, 95),
            Ring('Financial Security', cx, cy, 100, 60),
            Ring('Audit & Compliance', cx, cy, 45, 28),
        ],
    )
    specs.append(s)

    # ── 6.1 Production Infrastructure ──
    s = DiagramSpec(
        slug='14-Production-Infrastructure', figure='Figure 6.1',
        title='Production Infrastructure Architecture',
        purpose='Production deployment topology from internet edge through containers, business services, databases, and monitoring.',
        section='6. Infrastructure Architecture', insert_after='Section 6 Introduction',
        description='VPS production stack with twelve containers.',
        bands=layout_bands_stack([
            'Internet Edge', 'Load Balancer & Gateway', 'Application Containers',
            'Business Services', 'Data Tier', 'Monitoring & Observability',
        ], top, bottom),
    )
    rows6 = [
        [('Internet', 'ext', 'cloud'), ('Public DNS', 'ext', 'cloud'), ('Client Traffic', 'ext', 'user')],
        [('Reverse Proxy', 'infra', 'server'), ('TLS Gateway', 'infra', 'shield'), ('Route Control', 'infra', 'server')],
        [('Web Container', 'infra', 'container'), ('Admin Container', 'infra', 'container'), ('API Container', 'infra', 'container'), ('Engine Container', 'infra', 'container')],
        [('Trading Service', 'app', 'server'), ('Wallet Service', 'app', 'wallet'), ('Settlement Service', 'app', 'server'), ('Indexer Service', 'app', 'server')],
        [('Operational Database', 'data', 'database'), ('Distributed Cache', 'data', 'database'), ('Message Broker', 'data', 'queue'), ('Event Streaming Platform', 'data', 'queue')],
        [('Prometheus', 'mon', 'monitor'), ('Grafana', 'mon', 'monitor'), ('Health Checks', 'mon', 'monitor')],
    ]
    for band, row in zip(s.bands, rows6):
        s.components.extend(place_row_in_band(row, band))
    specs.append(s)

    # ── 6.2 Docker Deployment ──
    s = DiagramSpec(
        slug='15-Docker-Deployment', figure='Figure 6.2',
        title='Container Deployment Architecture',
        purpose='Docker Compose deployment showing application containers, data containers, network, and persistent volumes.',
        section='6. Infrastructure Architecture', insert_after='Figure 6.1',
        description='Container layout on exchange-production Docker network.',
        bands=layout_bands_stack(['Application Containers', 'Platform Containers', 'Persistent Volumes'], top, bottom),
    )
    for band, row in zip(s.bands, [
        [('Web Application', 'infra', 'container'), ('Administration Portal', 'infra', 'container'), ('API Server', 'infra', 'container'), ('Matching Engine', 'infra', 'container'), ('Indexer', 'infra', 'container')],
        [('Operational Database', 'infra', 'database'), ('Distributed Cache', 'infra', 'database'), ('Message Broker', 'infra', 'queue'), ('Event Streaming Platform', 'infra', 'queue'), ('Monitoring', 'infra', 'monitor')],
        [('Database Volume', 'data', 'database'), ('Cache Volume', 'data', 'database'), ('Engine Volume', 'data', 'database'), ('Queue Volume', 'data', 'database')],
    ]):
        s.components.extend(place_row_in_band(row, band))
    s.groups = [Group('Docker Network: exchange-production', MARGIN + 8, top + 8, DIAG_W - 16, bottom - top - 16)]
    specs.append(s)

    # ── 6.3 Network Topology ──
    s = DiagramSpec(
        slug='16-Network-Topology', figure='Figure 6.3',
        title='Network Topology Architecture',
        purpose='Public, private, and internal network segmentation for exchange services.',
        section='6. Infrastructure Architecture', insert_after='Figure 6.2',
        description='Network zones and service placement.',
        bands=layout_bands_stack(['Public Zone', 'Private Application Zone', 'Internal Services Zone'], top, bottom),
    )
    s.components.extend(place_row_in_band([('Internet Clients', 'ext', 'cloud'), ('Public Endpoint', 'infra', 'server'), ('Edge Firewall', 'infra', 'shield')], s.bands[0]))
    s.components.extend(place_row_in_band([('Web Tier', 'infra', 'container'), ('Admin Tier', 'infra', 'container'), ('API Tier', 'infra', 'container')], s.bands[1]))
    s.components.extend(place_row_in_band([('Matching Engine', 'infra', 'server'), ('Data Services', 'data', 'database'), ('Messaging Services', 'data', 'queue')], s.bands[2]))
    specs.append(s)

    # ── 6.4 Monitoring ──
    s = DiagramSpec(
        slug='17-Monitoring-Architecture', figure='Figure 6.4',
        title='Monitoring & Observability Architecture',
        purpose='Metrics collection, health monitoring, alerting, dashboards, and operational visibility.',
        section='6. Infrastructure Architecture', insert_after='Figure 6.3',
        description='Observability stack and health endpoints.',
        bands=layout_bands_stack(['Instrumentation Targets', 'Collection & Rules', 'Visualization & Response'], top, bottom),
    )
    for band, row in zip(s.bands, [
        [('API Platform', 'app', 'api'), ('Matching Engine', 'app', 'server'), ('Indexer', 'app', 'server'), ('Wallet Service', 'app', 'wallet')],
        [('Metrics Store', 'mon', 'monitor'), ('Alert Rules', 'mon', 'monitor'), ('Health Probes', 'mon', 'monitor')],
        [('Operations Dashboard', 'mon', 'monitor'), ('Admin Monitoring', 'mon', 'monitor'), ('Incident Response', 'mon', 'monitor')],
    ]):
        s.components.extend(place_row_in_band(row, band))
    specs.append(s)

    # ── 7.1 API Architecture ──
    s = DiagramSpec(
        slug='18-API-Architecture', figure='Figure 7.1',
        title='API Platform Architecture',
        purpose='How clients reach business and financial APIs through gateway, authentication, and data persistence layers.',
        section='7. Integration Architecture', insert_after='Section 7 Introduction',
        description='REST API platform with 447 endpoints across customer and admin surfaces.',
        bands=layout_bands_stack(['Client Tier', 'Gateway & Security', 'API Domains', 'Persistence'], top, bottom),
    )
    s.components.extend(place_row_in_band([('Web Client', 'ext', 'user'), ('Admin Client', 'ext', 'user'), ('Mobile Client', 'ext', 'user'), ('Partner API', 'ext', 'api')], s.bands[0]))
    s.components.extend(place_row_in_band([('API Gateway', 'infra', 'server'), ('Authentication', 'app', 'shield'), ('Rate Control', 'app', 'shield')], s.bands[1]))
    s.components.extend(place_row_in_band([('Core Business Services', 'app', 'api'), ('Financial APIs', 'app', 'wallet'), ('Admin APIs', 'app', 'api'), ('Public REST APIs', 'app', 'api')], s.bands[2]))
    s.components.extend(place_row_in_band([('Operational Database', 'data', 'database'), ('Session Cache', 'data', 'database')], s.bands[3]))
    specs.append(s)

    # ── 7.2 Request Lifecycle ──
    s = DiagramSpec(
        slug='19-Request-Lifecycle', figure='Figure 7.2',
        title='API Request Lifecycle',
        purpose='Lifecycle of an authenticated API request through gateway, business processing, optional settlement, and response.',
        section='7. Integration Architecture', insert_after='Figure 7.1',
        description='Request path for synchronous API operations.',
        pipeline=[
            ('Client Request', 'ext', 'api'), ('Gateway', 'infra', 'server'), ('Authentication', 'app', 'shield'),
            ('Business Service', 'app', 'server'), ('Financial Processing', 'app', 'wallet'), ('Response Delivery', 'app', 'api'),
        ],
        pipeline_x=MARGIN + 60, pipeline_y=centered_pipeline_y(6),
    )
    specs.append(s)

    # ── 8.1 Complete Ecosystem ──
    s = DiagramSpec(
        slug='20-Complete-Exchange-Ecosystem', figure='Figure 8.1',
        title='Complete Exchange Ecosystem',
        purpose='Full ecosystem map showing actors, channels, platform services, financial processing, infrastructure, and observability.',
        section='8. Enterprise Context', insert_after='Section 8 Introduction',
        description='Enterprise ecosystem view for due diligence and board presentation.',
        bands=layout_bands_stack([
            'Actors & Channels', 'Edge & Gateway', 'Application Services',
            'Domain Services', 'Financial Processing', 'Platform Infrastructure', 'Observability',
        ], top, bottom),
    )
    eco = [
        [('Retail Users', 'ext', 'user'), ('Institutions', 'ext', 'user'), ('Administrators', 'ext', 'user'), ('Web', 'app', 'server'), ('Mobile', 'app', 'server')],
        [('API Consumers', 'ext', 'api'), ('Gateway', 'infra', 'server'), ('Security Edge', 'infra', 'shield')],
        [('Customer Platform', 'app', 'server'), ('Admin Platform', 'app', 'server'), ('Integration Platform', 'app', 'api')],
        [('Trading', 'app', 'server'), ('Wallet', 'app', 'wallet'), ('Treasury', 'app', 'wallet'), ('P2P', 'app', 'server'), ('Compliance', 'app', 'shield')],
        [('Notification', 'app', 'queue'), ('Settlement', 'app', 'server'), ('Ledger', 'data', 'database')],
        [('Database', 'data', 'database'), ('Distributed Cache', 'data', 'database'), ('Message Broker', 'data', 'queue'), ('Containers', 'infra', 'container')],
        [('Monitoring', 'mon', 'monitor'), ('Dashboards', 'mon', 'monitor'), ('Health Management', 'mon', 'monitor')],
    ]
    for band, row in zip(s.bands, eco):
        s.components.extend(place_row_in_band(row, band))
    specs.append(s)

    return specs


def write_readme(specs: list[DiagramSpec]) -> None:
    lines = [
        '# Enterprise Architecture Diagram Pack v4',
        '',
        '**Standard:** Microsoft Azure Architecture Center / AWS Well-Architected / Oracle Architecture Reference quality',
        '',
        '## Design Standards',
        '',
        '| Element | Specification |',
        '| --- | --- |',
        '| Application Services | Blue (#DEECF9 / #0078D4) |',
        '| Data Stores | Yellow (#FFF4CE / #FFB900) |',
        '| Monitoring | Green (#D4EDDA / #107C10) |',
        '| Infrastructure | Grey (#F3F2F1 / #605E5C) |',
        '| External Systems | Purple (#E7DFEC / #5C2D91) |',
        '| Typography | Segoe UI — Title 28pt, Purpose 14pt, Layer 16pt, Component 15pt, Legend 11pt |',
        '| Box design | 12px radius, soft shadow, 2px stroke, 24px padding, icon left |',
        '| Connectors | Primary solid blue 2.5px; secondary grey dashed 1px; orthogonal routing |',
        '| Layout | ~80–85% printable area; compact footer legend outside drawing zone |',
        '',
        'Each figure includes: Figure number, Title, Purpose, Layer labels, Footer legend, Fluent-style icons, Grouped domains.',
        '',
        '## Figure Index',
        '',
        '| Figure | File Prefix | Section | Insert After | Purpose |',
        '| --- | --- | --- | --- | --- |',
    ]
    for s in specs:
        lines.append(f'| {s.figure} | `{s.slug}` | {s.section} | {s.insert_after} | {s.purpose} |')
    (OUT / 'README.md').write_text('\n'.join(lines) + '\n')


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    specs = build_all()
    for spec in specs:
        base = OUT / spec.slug
        svg = base.with_suffix('.svg')
        svg.write_text(render_svg(spec))
        base.with_suffix('.drawio').write_text(render_drawio(spec))
        export_png(svg, base.with_suffix('.png'))
        export_pdf(svg, base.with_suffix('.pdf'))
        render_pptx(spec, base.with_suffix('.pptx'))
        print(f'OK {spec.slug}')
    write_readme(specs)
    print(f'Enterprise pack v4 complete: {len(specs)} figures')


if __name__ == '__main__':
    main()
